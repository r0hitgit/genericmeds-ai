#!/usr/bin/env python3
"""
Comprehensive Excel Master Ingestion Script for GenericMeds AI.
Ingests:
1. Verified Government & Market Data Sources
2. 10 Hospital Jan Aushadhi Kendras
3. 15 Active-Ingredient Safe Molecule Guidelines & Dispensing Rules
4. 39 Flagship Branded Reference Products with Market MRPs
5. All 2,439 PMBI Reconciled Master Catalogue Products across all 68 Therapeutic Groups:
   - 2,052 Active Priced Items
   - 387 Under-Process Items (Stored strictly as amount_inr=None, verification_status='under_process')
"""

import os
import sys
import re
import openpyxl
from datetime import datetime, timezone
from sqlalchemy import func

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.core.database import SessionLocal, init_db
from app.models.source import DataSource
from app.models.ingredient import ActiveIngredient, ProductIngredient
from app.models.medicine import MedicineProduct
from app.models.pricing import MedicinePrice
from app.models.kendra import JanAushadhiKendra
from app.models.guideline import MoleculeGuideline
from app.services.source_verification import SourceVerificationService
from scripts.process_all_pmbjp_data import (
    clean_number,
    parse_pack_info,
    parse_dosage_form,
    parse_release_type,
    parse_prescription_status,
)
from scripts.build_datasets import FLAGSHIP_COMPARISONS


EXCEL_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data", "imports", "pmbjp_master_2439.xlsx"))


def parse_ingredients_advanced(text: str):
    """
    Parses active ingredients and strengths with high clinical precision.
    Removes standard pharmacopeia tokens (IP, BP, USP) and trailing dosage form labels.
    """
    if not text:
        return []
    s = re.sub(r"\b(IP|BP|USP)\b", "", text, flags=re.IGNORECASE).strip()
    chunks = re.split(r"\s*\+\s*|\s+and\s+|\s*&\s*|,\s*", s)
    results = []

    pattern = re.compile(
        r"^(?P<before>.*?)(?:\s*(?P<val>\d+(?:\.\d+)?)\s*(?P<unit>mg|g|gm|mcg|ml|iu|%)\b)(?P<after>.*)$",
        re.IGNORECASE,
    )
    dosage_words = re.compile(
        r"\b(Tablets|Tablet|Capsules|Capsule|Injection|Inj|Syrup|Suspension|Ointment|Cream|Gel|Drops|Solution|Powder|Respules|Respule|Lotion|Mouthwash|with WFI|Tab|Cap)\b",
        re.IGNORECASE,
    )

    for chunk in chunks:
        chunk = chunk.strip(" ()-+")
        if not chunk:
            continue
        m = pattern.search(chunk)
        if m:
            before = dosage_words.sub("", m.group("before")).strip(" ()-+")
            after = dosage_words.sub("", m.group("after")).strip(" ()-+")
            name = f"{before} {after}".strip()
            if not name:
                name = m.group("before").strip()
            val = float(m.group("val"))
            unit = m.group("unit").lower()
            if unit == "gm":
                unit = "g"
            results.append({"name": name, "strength_value": val, "strength_unit": unit})
        else:
            clean_name = dosage_words.sub("", chunk).strip(" ()-+")
            results.append({"name": clean_name or chunk, "strength_value": 0.0, "strength_unit": "mg"})

    return results


def run_master_ingestion():
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Master Excel file not found at: {EXCEL_PATH}")

    print(f"Loading master workbook from: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    init_db()
    db = SessionLocal()
    now = datetime.now(timezone.utc)

    try:
        # 1. Sources
        print("Step 1: Verifying default data sources...")
        SourceVerificationService.ensure_default_sources(db)
        src_jan = db.query(DataSource).filter(DataSource.source_name == "PMBI Jan Aushadhi (PMBJP)").first()
        src_market = db.query(DataSource).filter(DataSource.source_name == "Verified Branded Market Retail").first()

        # 2. Hospital Kendras
        print("Step 2: Ingesting Hospital Jan Aushadhi Kendras...")
        ws_ken = wb["Hospital Kendras & PMBI Rules"]
        k_created = 0
        k_updated = 0
        for r in list(ws_ken.iter_rows(values_only=True))[3:]:
            k_code = str(r[0]).strip() if r[0] else None
            if not k_code:
                continue
            k_name = str(r[1]).strip() if r[1] else ""
            k_city = str(r[2]).strip() if r[2] else ""
            k_loc = str(r[3]).strip() if r[3] else ""
            k_hours = str(r[4]).strip() if r[4] else ""
            k_scope = str(r[5]).strip() if r[5] else ""
            k_rules = str(r[6]).strip() if r[6] else ""
            k_help = str(r[7]).strip() if r[7] else ""
            is_247 = "24x7" in k_hours.lower()

            ex_kendra = db.query(JanAushadhiKendra).filter(JanAushadhiKendra.kendra_code == k_code).first()
            if ex_kendra:
                ex_kendra.hospital_name = k_name
                ex_kendra.city_state = k_city
                ex_kendra.location = k_loc
                ex_kendra.operating_hours = k_hours
                ex_kendra.portfolio_scope = k_scope
                ex_kendra.dispensing_rules = k_rules
                ex_kendra.helpline_url = k_help
                ex_kendra.is_24x7 = is_247
                k_updated += 1
            else:
                db.add(JanAushadhiKendra(
                    kendra_code=k_code,
                    hospital_name=k_name,
                    city_state=k_city,
                    location=k_loc,
                    operating_hours=k_hours,
                    portfolio_scope=k_scope,
                    dispensing_rules=k_rules,
                    helpline_url=k_help,
                    is_24x7=is_247,
                ))
                k_created += 1
        db.commit()
        print(f"   -> Kendras: {k_created} created, {k_updated} updated.")

        # 3. Molecule Guidelines
        print("Step 3: Ingesting Safe Molecule Guidelines & Dispensing Rules...")
        ws_mol = wb["Molecule Price Lookup (Safe)"]
        m_created = 0
        m_updated = 0
        for r in list(ws_mol.iter_rows(values_only=True))[3:]:
            m_id = str(r[0]).strip() if r[0] else None
            if not m_id:
                continue
            m_grp = str(r[1]).strip() if r[1] else ""
            m_ing = str(r[2]).strip() if r[2] else ""
            m_forms = str(r[3]).strip() if r[3] else ""
            m_codes = str(r[4]).strip() if r[4] else ""
            m_packs = str(r[5]).strip() if r[5] else ""
            m_pmbi_mrp = str(r[6]).strip() if r[6] else ""
            m_brand_mrp = str(r[7]).strip() if r[7] else ""
            m_diff = str(r[8]).strip() if r[8] else ""
            m_rule = str(r[9]).strip() if r[9] else ""
            m_url = str(r[10]).strip() if r[10] else ""

            ex_mol = db.query(MoleculeGuideline).filter(MoleculeGuideline.lookup_id == m_id).first()
            if ex_mol:
                ex_mol.therapeutic_class = m_grp
                ex_mol.active_ingredient = m_ing
                ex_mol.strengths_forms = m_forms
                ex_mol.pmbi_codes = m_codes
                ex_mol.pack_sizes = m_packs
                ex_mol.pmbi_mrp = m_pmbi_mrp
                ex_mol.branded_mrp = m_brand_mrp
                ex_mol.price_diff_pct = m_diff
                ex_mol.dispensing_rule = m_rule
                ex_mol.reference_url = m_url
                m_updated += 1
            else:
                db.add(MoleculeGuideline(
                    lookup_id=m_id,
                    therapeutic_class=m_grp,
                    active_ingredient=m_ing,
                    strengths_forms=m_forms,
                    pmbi_codes=m_codes,
                    pack_sizes=m_packs,
                    pmbi_mrp=m_pmbi_mrp,
                    branded_mrp=m_brand_mrp,
                    price_diff_pct=m_diff,
                    dispensing_rule=m_rule,
                    reference_url=m_url,
                ))
                m_created += 1
        db.commit()
        print(f"   -> Molecule Guidelines: {m_created} created, {m_updated} updated.")

        # 4. Flagship Branded References
        print("Step 4: Ingesting 39 Flagship Branded References with Market MRPs...")
        b_created = 0
        b_updated = 0
        for item in FLAGSHIP_COMPARISONS:
            b_name = item["brand_ref"].split("/")[0].strip()
            g_name = item["name"]
            p_size, p_unit = parse_pack_info(item["pack"])
            d_form, a_route = parse_dosage_form(g_name, item["category"])
            r_type = parse_release_type(g_name)
            rx_status = parse_prescription_status(item["category"], g_name)

            prod = db.query(MedicineProduct).filter(
                MedicineProduct.brand_name.ilike(b_name),
                MedicineProduct.is_jan_aushadhi == False,
            ).first()

            if prod:
                prod.generic_name = g_name
                prod.pack_size = p_size
                prod.pack_unit = p_unit
                prod.updated_at = now
                b_updated += 1
            else:
                prod = MedicineProduct(
                    brand_name=b_name,
                    generic_name=g_name,
                    manufacturer_name="Market Reference Manufacturer",
                    dosage_form=d_form,
                    administration_route=a_route,
                    release_type=r_type,
                    pack_size=p_size,
                    pack_unit=p_unit,
                    prescription_status=rx_status,
                    is_jan_aushadhi=False,
                    therapeutic_group=item["category"],
                    source_id=src_market.id,
                    created_at=now,
                    updated_at=now,
                )
                db.add(prod)
                db.flush()
                b_created += 1

                # Ingredients
                parsed_ings = parse_ingredients_advanced(g_name)
                seen_flagship_ings = set()
                for ing in parsed_ings:
                    disp = ing["name"].strip().title()
                    norm = ing["name"].strip().lower()
                    if not norm or norm in seen_flagship_ings:
                        continue
                    seen_flagship_ings.add(norm)
                    act = db.query(ActiveIngredient).filter(ActiveIngredient.normalized_name == norm).first()
                    if not act:
                        act = ActiveIngredient(normalized_name=norm, display_name=disp)
                        db.add(act)
                        db.flush()
                    pi = ProductIngredient(
                        product_id=prod.id,
                        ingredient_id=act.id,
                        strength_value=ing["strength_value"],
                        strength_unit=ing["strength_unit"],
                    )
                    db.add(pi)

            # Price record
            mrp_val = item["market_price"]
            ex_price = db.query(MedicinePrice).filter(
                MedicinePrice.product_id == prod.id,
                MedicinePrice.price_type == "MRP",
            ).first()
            if ex_price:
                ex_price.amount_inr = mrp_val
                ex_price.pack_quantity = p_size
                ex_price.unit_of_measure = p_unit
                ex_price.verified_at = now
                ex_price.verification_status = "verified"
            else:
                db.add(MedicinePrice(
                    product_id=prod.id,
                    amount_inr=mrp_val,
                    price_type="MRP",
                    pack_quantity=p_size,
                    unit_of_measure=p_unit,
                    source_id=src_market.id,
                    source_url="https://nppaipdms.gov.in/NPPA/PharmaSahiDaam/searchMedicine",
                    verified_at=now,
                    verification_status="verified",
                ))
        db.commit()
        print(f"   -> Flagship Branded References: {b_created} created, {b_updated} updated.")

        # 5. Ingest Complete 2,439 PMBI Master Basket
        print("Step 5: Ingesting complete 2,439 PMBI Master Catalogue from Excel...")
        ws_cat = wb["PMBI Master Catalog (2439)"]
        p_created = 0
        p_updated = 0
        p_priced = 0
        p_unpriced = 0

        # Cache known active ingredients for fast lookup
        known_ingredients = {ing.normalized_name: ing.id for ing in db.query(ActiveIngredient).all()}

        for r in list(ws_cat.iter_rows(values_only=True))[3:]:
            code = str(r[0]).strip() if r[0] else None
            raw_name = str(r[1]).strip() if r[1] else None
            pack_raw = str(r[2]).strip() if r[2] else ""
            mrp_raw = r[3]
            status_raw = str(r[4]).strip() if r[4] else ""
            group_raw = str(r[5]).strip() if r[5] else ""
            source_url = str(r[8]).strip() if r[8] else "https://janaushadhi.gov.in/"

            if not code or not raw_name:
                continue

            # Standardized brand name for Jan Aushadhi
            brand_name = raw_name if raw_name.lower().startswith("jan") else f"Jan Aushadhi {raw_name}"
            brand_name = brand_name[:250]
            generic_name = raw_name[:490]

            pack_size, pack_unit = parse_pack_info(pack_raw)
            d_form, a_route = parse_dosage_form(raw_name, group_raw)
            r_type = parse_release_type(raw_name)
            rx_status = parse_prescription_status(group_raw, raw_name)
            price_val = clean_number(mrp_raw)

            # Query product by PMBI code or exact brand name
            prod = db.query(MedicineProduct).filter(
                (MedicineProduct.pmbi_code == code) |
                ((MedicineProduct.brand_name == brand_name) & (MedicineProduct.pack_size == pack_size))
            ).first()

            if prod:
                prod.generic_name = generic_name
                prod.pmbi_code = code
                prod.therapeutic_group = group_raw
                prod.dosage_form = d_form
                prod.administration_route = a_route
                prod.release_type = r_type
                prod.pack_size = pack_size
                prod.pack_unit = pack_unit
                prod.is_jan_aushadhi = True
                prod.updated_at = now
                p_updated += 1
            else:
                prod = MedicineProduct(
                    brand_name=brand_name,
                    generic_name=generic_name,
                    manufacturer_name="PMBI Jan Aushadhi",
                    dosage_form=d_form,
                    administration_route=a_route,
                    release_type=r_type,
                    pack_size=pack_size,
                    pack_unit=pack_unit,
                    prescription_status=rx_status,
                    is_jan_aushadhi=True,
                    pmbi_code=code,
                    therapeutic_group=group_raw,
                    source_id=src_jan.id,
                    created_at=now,
                    updated_at=now,
                )
                db.add(prod)
                db.flush()
                p_created += 1

                # Parse ingredients
                parsed_ings = parse_ingredients_advanced(raw_name)
                seen_norm_ings = set()
                for ing in parsed_ings:
                    norm = ing["name"].strip().lower()
                    disp = ing["name"].strip().title()
                    if not norm or norm in seen_norm_ings:
                        continue
                    seen_norm_ings.add(norm)

                    if norm in known_ingredients:
                        ing_id = known_ingredients[norm]
                    else:
                        act = ActiveIngredient(normalized_name=norm, display_name=disp)
                        db.add(act)
                        db.flush()
                        ing_id = act.id
                        known_ingredients[norm] = ing_id

                    pi = ProductIngredient(
                        product_id=prod.id,
                        ingredient_id=ing_id,
                        strength_value=ing["strength_value"],
                        strength_unit=ing["strength_unit"],
                    )
                    db.add(pi)

            # Price handling
            ex_price = db.query(MedicinePrice).filter(
                MedicinePrice.product_id == prod.id,
                MedicinePrice.price_type == "government_listed_price",
            ).first()

            if price_val is not None and price_val > 0:
                p_priced += 1
                if ex_price:
                    ex_price.amount_inr = price_val
                    ex_price.pack_quantity = pack_size
                    ex_price.unit_of_measure = pack_unit
                    ex_price.verified_at = now
                    ex_price.verification_status = "verified"
                else:
                    db.add(MedicinePrice(
                        product_id=prod.id,
                        amount_inr=price_val,
                        price_type="government_listed_price",
                        pack_quantity=pack_size,
                        unit_of_measure=pack_unit,
                        source_id=src_jan.id,
                        source_url=source_url,
                        verified_at=now,
                        verification_status="verified",
                    ))
            else:
                p_unpriced += 1
                if ex_price:
                    ex_price.amount_inr = None
                    ex_price.pack_quantity = pack_size
                    ex_price.unit_of_measure = pack_unit
                    ex_price.verified_at = now
                    ex_price.verification_status = "under_process"
                else:
                    db.add(MedicinePrice(
                        product_id=prod.id,
                        amount_inr=None,
                        price_type="government_listed_price",
                        pack_quantity=pack_size,
                        unit_of_measure=pack_unit,
                        source_id=src_jan.id,
                        source_url=source_url,
                        verified_at=now,
                        verification_status="under_process",
                    ))

            # Commit periodically every 250 records
            if (p_created + p_updated) % 250 == 0:
                db.commit()

        src_jan.last_checked_at = now
        src_jan.last_successful_sync_at = now
        db.commit()

        print(f"   -> PMBI Master Basket: {p_created} created, {p_updated} updated.")
        print(f"   -> Priced Items: {p_priced}, Under-Process Items (No MRP): {p_unpriced}")
        print("Excel Master Integration Completed Successfully!")

    finally:
        db.close()


if __name__ == "__main__":
    run_master_ingestion()
